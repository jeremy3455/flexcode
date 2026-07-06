import json
import os
from datetime import datetime

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "documentos_generados")


def _ensure_output_dir():
    os.makedirs(OUTPUT_DIR, exist_ok=True)


def _safe_filename(name: str, ext: str) -> str:
    safe = "".join(c if c.isalnum() or c in " _-" else "_" for c in name)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{safe}_{ts}{ext}"


def generate_pdf(title: str, content: str, filename: str = "") -> str:
    from fpdf import FPDF

    _ensure_output_dir()
    path = os.path.join(OUTPUT_DIR, _safe_filename(filename or title, ".pdf"))

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 12, title, new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(6)
    pdf.set_font("Helvetica", "", 11)

    for paragraph in content.split("\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            pdf.ln(4)
            continue
        try:
            pdf.multi_cell(0, 6, paragraph)
        except Exception:
            pdf.multi_cell(0, 6, paragraph.encode("latin-1", "replace").decode("latin-1"))

    pdf.output(path)
    fname = os.path.basename(path)
    return f"PDF generado: [{fname}](/descargar/{fname})"


def generate_docx(title: str, content: str, filename: str = "") -> str:
    from docx import Document
    from docx.shared import Pt

    _ensure_output_dir()
    path = os.path.join(OUTPUT_DIR, _safe_filename(filename or title, ".docx"))

    doc = Document()
    doc.add_heading(title, 0)

    for paragraph in content.split("\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        p = doc.add_paragraph(paragraph)
        for run in p.runs:
            run.font.size = Pt(11)

    doc.save(path)
    fname = os.path.basename(path)
    return f"Word generado: [{fname}](/descargar/{fname})"


def generate_xlsx(headers: list, rows: list, filename: str = "") -> str:
    from openpyxl import Workbook

    _ensure_output_dir()
    safe_name = filename or "datos"
    path = os.path.join(OUTPUT_DIR, _safe_filename(safe_name, ".xlsx"))

    wb = Workbook()
    ws = wb.active
    ws.title = safe_name[:31]

    if headers:
        ws.append(headers)
        from openpyxl.styles import Font
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=1, column=col_idx).font = Font(bold=True)

    for row in rows:
        ws.append(row)

    wb.save(path)
    fname = os.path.basename(path)
    return f"Excel generado: [{fname}](/descargar/{fname})"
